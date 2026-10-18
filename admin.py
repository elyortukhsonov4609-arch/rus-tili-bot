"""Admin panel: arizalar, o'quvchilar, vaqt, natijalar, darslar, testlar, guruh sozlamalari."""
import html
import io
import re
import time
from datetime import datetime

from aiogram import Bot, F, Router
from aiogram.filters import Command, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import BufferedInputFile, CallbackQuery, Message
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import db
from common import MENU_BUTTONS, BTN_ADMIN, BTN_LESSONS, BTN_RATING, BTN_RESULTS, BTN_TESTS, IsAdmin, PrivateChat, chunk_text, fmt_dur, kb, main_menu, pct, safe_answer, safe_edit
from config import DB_PATH, INTERVAL_CHOICES, PER_ATTEMPT_CHOICES, SECONDS_CHOICES, TZ
from docx_tools import export_tests, lesson_blocks, paginate, parse_tests
from group import kick_scheduler, count_buttons, group_top_text, post_question, running, start_group_test, stop_group_test

router = Router(name="admin")
router.message.filter(PrivateChat(), IsAdmin())
router.callback_query.filter(IsAdmin())


class AdmSt(StatesGroup):
    lesson_title = State()      # yangi dars nomi
    lesson_content = State()    # dars matni / fayli kutilmoqda
    test_file = State()         # yangi test uchun Word fayl
    test_title = State()        # yangi test nomi
    test_replace = State()      # mavjud testni Word fayl bilan almashtirish
    time_input = State()        # vaqt kiritish (kunlik vazifa, eslatma, hisobot, guruh jadvali)


MENU_TEXTS = MENU_BUTTONS


async def _left_input(msg: Message, state: FSMContext) -> bool:
    """Admin kiritish jarayonida menyu tugmasi yoki buyruq bossa — jarayonni yakunlaydi."""
    if msg.text in MENU_TEXTS or msg.text.startswith("/"):
        await state.clear()
        await msg.answer("ℹ️ Kiritish yakunlandi. Kerakli tugmani qayta bosing.", reply_markup=main_menu(msg.from_user.id))
        return True
    return False


def yn(v) -> str:
    return "✅" if int(v) else "❌"


# ======================= Bosh menyu =======================
async def home_markup():
    pending = (await db.fetchone("SELECT COUNT(*) AS n FROM users WHERE status='pending'"))["n"]
    return kb([
        [(f"🆕 Arizalar ({pending})", "adm:pending"), ("👥 O'quvchilar", "adm:users:0")],
        [("⏱ Vaqt sozlamalari", "adm:time"), ("📊 Natijalar", "adm:res")],
        [("📚 Darslar", "adm:les"), ("📝 Testlar", "adm:tst")],
        [("📅 Kunlik vazifa", "adm:dly"), ("👨‍👩‍👦 Ota-onalar", "adm:par")],
        [("👥 Guruh sozlamalari", "adm:grp"), ("⚙️ Boshqa", "adm:opt")],
    ])


HOME_TEXT = "⚙️ <b>Admin panel</b>\nKerakli bo'limni tanlang:"


@router.message(F.text == BTN_ADMIN)
@router.message(Command("admin"))
async def admin_home(msg: Message, state: FSMContext):
    await state.clear()
    await msg.answer(HOME_TEXT, reply_markup=await home_markup())


@router.callback_query(F.data == "adm:home")
async def admin_home_cb(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    await safe_edit(cb, HOME_TEXT, await home_markup())
    await cb.answer()


@router.message(Command("cancel"))
async def cancel(msg: Message, state: FSMContext):
    await state.clear()
    await msg.answer("Bekor qilindi.", reply_markup=main_menu(msg.from_user.id))


BACK = [("⬅️ Admin panel", "adm:home")]


# ======================= Arizalar va o'quvchilar =======================
async def group_buttons(prefix: str, current=None) -> list:
    """Har bir o'quv guruhi uchun tugma: [(nom, prefix:id)]."""
    return [[(("🔘 " if g["id"] == current else "") + g["name"], f"{prefix}:{g['id']}") for g in await db.study_groups()]]


@router.callback_query(F.data == "adm:pending")
async def pending(cb: CallbackQuery):
    rows = await db.fetchall("SELECT * FROM users WHERE status='pending' ORDER BY created_at")
    if not rows:
        await safe_edit(cb, "🆕 Yangi arizalar yo'q.", kb([BACK]))
        return await cb.answer()
    groups = await db.study_groups()
    btns = []
    for u in rows[:30]:
        btns.append([(f"👤 {u['full_name'][:28]}", "noop")])
        btns.append([(f"✅ {g['name']}", f"adm:ok:{u['id']}:{g['id']}") for g in groups] + [("❌", f"adm:no:{u['id']}")])
    btns.append([(f"✅ Hammasi → {g['name']}", f"adm:okall:{g['id']}") for g in groups])
    btns.append(BACK)
    await safe_edit(cb, f"🆕 <b>Arizalar: {len(rows)} ta</b>\nHar bir o'quvchini o'z guruhiga tasdiqlang:", kb(btns))
    await cb.answer()


async def _approve(bot: Bot, uid: int, gid: int):
    await db.execute("UPDATE users SET group_id=? WHERE id=?", (gid, uid))
    await db.set_user_status(uid, "approved")
    try:
        await bot.send_message(uid, f"🎉 O'qituvchi sizga ruxsat berdi! Guruhingiz: <b>{html.escape(await db.group_name(gid))}</b>\n"
                                    "Avval «📚 Darslar» bo'limida mavzuni o'qing, keyin test ishlang.\n"
                                    "Har kuni «📅 Bugungi vazifa»ni bajarib, 🔥 seriyangizni oshiring!",
                               reply_markup=main_menu(uid))
    except Exception:  # noqa: BLE001
        pass


@router.callback_query(F.data.startswith("adm:ok:"))
async def approve(cb: CallbackQuery, bot: Bot):
    parts = cb.data.split(":")
    uid = int(parts[2])
    u = await db.get_user(uid)
    if not u:
        return await cb.answer("Foydalanuvchi topilmadi", show_alert=True)
    if len(parts) < 4:   # guruh tanlanmagan (eski tugma) — so'raymiz
        await safe_edit(cb, f"👤 <b>{html.escape(u['full_name'])}</b> — qaysi guruhga qo'shilsin?",
                        kb(await group_buttons(f"adm:ok:{uid}") + [[("⬅️ Orqaga", "adm:pending")]]))
        return await cb.answer()
    gid = int(parts[3])
    await _approve(bot, uid, gid)
    gname = await db.group_name(gid)
    await cb.answer(f"✅ {u['full_name']} → {gname}")
    if cb.message.text and cb.message.text.startswith("🆕 Yangi ariza"):
        await safe_edit(cb, cb.message.html_text + f"\n\n✅ <b>Ruxsat berildi: {html.escape(gname)}</b>")
    else:
        await pending(cb)


@router.callback_query(F.data.startswith("adm:okall:"))
async def approve_all(cb: CallbackQuery, bot: Bot):
    gid = int(cb.data.split(":")[2])
    rows = await db.fetchall("SELECT id FROM users WHERE status='pending'")
    for r in rows:
        await _approve(bot, r["id"], gid)
    await cb.answer(f"✅ {len(rows)} ta o'quvchi → {await db.group_name(gid)}", show_alert=True)
    await pending(cb)


@router.callback_query(F.data.startswith("adm:no:"))
async def reject(cb: CallbackQuery, bot: Bot):
    uid = int(cb.data.split(":")[2])
    u = await db.get_user(uid)
    if not u:
        return await cb.answer("Foydalanuvchi topilmadi", show_alert=True)
    await db.set_user_status(uid, "rejected")
    try:
        await bot.send_message(uid, "❌ Afsuski, arizangiz rad etildi. O'qituvchingiz bilan bog'laning.")
    except Exception:  # noqa: BLE001
        pass
    await cb.answer(f"❌ {u['full_name']} — rad etildi")
    if cb.message.text and cb.message.text.startswith("🆕 Yangi ariza"):
        await safe_edit(cb, cb.message.html_text + "\n\n❌ <b>Rad etildi</b>")
    else:
        await pending(cb)


@router.callback_query(F.data.startswith("adm:users:"))
async def users(cb: CallbackQuery):
    parts = cb.data.split(":")
    page = int(parts[2])
    gid = int(parts[3]) if len(parts) > 3 else 0      # 0 = hamma guruhlar
    per = 20
    where, args = ("status!='pending'", ())
    if gid:
        where, args = ("status!='pending' AND group_id=?", (gid,))
    total = (await db.fetchone(f"SELECT COUNT(*) AS n FROM users WHERE {where}", args))["n"]
    rows = await db.fetchall(f"SELECT * FROM users WHERE {where} ORDER BY status, group_id, full_name LIMIT ? OFFSET ?",
                             args + (per, page * per))
    icon = {"approved": "🟢", "blocked": "🚫", "rejected": "⚪️"}
    short = {g["id"]: g["name"].split()[0] for g in await db.study_groups()}
    btns = [[(("🔘 " if gid == 0 else "") + "Hammasi", "adm:users:0:0")] +
            [(("🔘 " if gid == g["id"] else "") + g["name"], f"adm:users:0:{g['id']}") for g in await db.study_groups()]]
    btns += [[(f"{icon.get(u['status'], '')} {u['full_name']}" + (f" · {short.get(u['group_id'], '')}" if not gid else ""),
               f"adm:user:{u['id']}")] for u in rows]
    nav = []
    if page > 0:
        nav.append(("⬅️", f"adm:users:{page - 1}:{gid}"))
    if (page + 1) * per < total:
        nav.append(("➡️", f"adm:users:{page + 1}:{gid}"))
    if nav:
        btns.append(nav)
    btns.append(BACK)
    await safe_edit(cb, f"👥 <b>O'quvchilar: {total}</b>\n🟢 faol · 🚫 chiqarilgan · ⚪️ rad etilgan", kb(btns))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:user:"))
async def user_card_cb(cb: CallbackQuery):
    await user_card(cb, int(cb.data.split(":")[2]))


async def user_card(cb: CallbackQuery, uid: int):
    from daily import streak
    u = await db.get_user(uid)
    if not u:
        return await cb.answer("Topilmadi", show_alert=True)
    att = await db.fetchall("SELECT a.*, t.title FROM attempts a JOIN tests t ON t.id=a.test_id "
                            "WHERE a.user_id=? AND a.status='done' ORDER BY a.finished_at", (uid,))
    reads = await db.fetchall("SELECT l.num FROM lesson_reads r JOIN lessons l ON l.id=r.lesson_id "
                              "WHERE r.user_id=? ORDER BY l.num", (uid,))
    par = await db.fetchall("SELECT * FROM parents WHERE child_id=?", (uid,))
    lines = [f"👤 <b>{html.escape(u['full_name'])}</b>",
             f"👥 Guruh: <b>{html.escape(await db.group_name(u['group_id']))}</b>",
             f"🔗 @{u['username']}" if u["username"] else "🔗 username yo'q",
             f"🆔 <code>{uid}</code>",
             f"🔥 Seriya: {await streak(uid)} kun",
             f"👨‍👩‍👦 Ota-ona: {', '.join(html.escape(p['name'] or '—') for p in par) or 'ulanmagan'}",
             f"📚 O'qilgan darslar: {', '.join(str(r['num']) for r in reads) or '—'}",
             f"📝 Ishlangan testlar: {len(att)}", ""]
    for a in att[-25:]:
        d = datetime.fromtimestamp(a["finished_at"], TZ).strftime("%d.%m %H:%M")
        lines.append(f"• {html.escape(a['title'])}: <b>{a['correct']}/{a['total']}</b> "
                     f"({pct(a['correct'], a['total'])}%) · ⏱ {fmt_dur(a['duration'])} · {d}")
    btns = [[("👥 Guruhini o'zgartirish", f"adm:ugrp:{uid}")]]
    if u["status"] == "approved":
        btns.append([("🚫 Botdan chiqarish", f"adm:block:{uid}")])
    else:
        btns.append([("✅ Ruxsat berish", f"adm:ok:{uid}")])
    for p in par:
        btns.append([(f"❌ Ota-onani uzish: {(p['name'] or '—')[:20]}", f"adm:punl:{p['parent_id']}:{uid}")])
    btns.append([("📊 Haftalik hisobotni ko'rish", f"adm:prep:{uid}")])
    btns.append([("⬅️ O'quvchilar", "adm:users:0")])
    text = "\n".join(lines)
    await safe_edit(cb, text[:4000], kb(btns))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:ugrp:"))
async def user_group_change(cb: CallbackQuery):
    parts = cb.data.split(":")
    uid = int(parts[2])
    if len(parts) == 4:
        await db.execute("UPDATE users SET group_id=? WHERE id=?", (int(parts[3]), uid))
        await cb.answer("Saqlandi")
        return await user_card(cb, uid)
    u = await db.get_user(uid)
    await safe_edit(cb, f"👤 <b>{html.escape(u['full_name'])}</b> — yangi guruhni tanlang:",
                    kb(await group_buttons(f"adm:ugrp:{uid}", u["group_id"]) + [[("⬅️ Orqaga", f"adm:user:{uid}")]]))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:punl:"))
async def parent_unlink(cb: CallbackQuery):
    _, _, pid, cid = cb.data.split(":")
    await db.execute("DELETE FROM parents WHERE parent_id=? AND child_id=?", (int(pid), int(cid)))
    await cb.answer("Ota-ona uzildi", show_alert=True)
    if cb.message.text and cb.message.text.startswith("👨‍👩‍👦"):
        await safe_edit(cb, cb.message.html_text + "\n\n❌ <b>Uzildi</b>")
    else:
        await user_card(cb, int(cid))


@router.callback_query(F.data.startswith("adm:prep:"))
async def parent_report_preview(cb: CallbackQuery):
    from parents import build_report
    await cb.message.answer(await build_report(int(cb.data.split(":")[2])))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:block:"))
async def block(cb: CallbackQuery):
    uid = int(cb.data.split(":")[2])
    await db.set_user_status(uid, "blocked")
    await cb.answer("🚫 Chiqarildi. Endi u botdan foydalana olmaydi.", show_alert=True)
    await user_card(cb, uid)


# ======================= Vaqt sozlamalari =======================
@router.callback_query(F.data == "adm:time")
async def time_menu(cb: CallbackQuery):
    ps, gs = await db.get_int("private_seconds"), await db.get_int("group_seconds")
    row1 = [(("🔘 " if s == ps else "") + f"{s} s", f"adm:ps:{s}") for s in SECONDS_CHOICES]
    row2 = [(("🔘 " if s == gs else "") + f"{s} s", f"adm:gs:{s}") for s in SECONDS_CHOICES]
    text = ("⏱ <b>Bitta savolga beriladigan vaqt</b>\n\n"
            f"👤 Shaxsiy testda (bot ichida): <b>{ps} soniya</b>\n"
            f"👥 Guruhda: <b>{gs} soniya</b>\n\n"
            "Yuqori qator — shaxsiy test, pastki qator — guruh.")
    await safe_edit(cb, text, kb([[("👤 Shaxsiy test:", "noop")], row1, [("👥 Guruh:", "noop")], row2, BACK]))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:ps:") | F.data.startswith("adm:gs:"))
async def time_set(cb: CallbackQuery):
    _, kind, sec = cb.data.split(":")
    await db.set_setting("private_seconds" if kind == "ps" else "group_seconds", int(sec))
    await cb.answer(f"✅ {sec} soniya saqlandi")
    await time_menu(cb)


# ======================= Natijalar =======================
@router.callback_query(F.data == "adm:res")
async def results_menu(cb: CallbackQuery):
    tests = await db.fetchall(
        "SELECT t.id, t.title, t.group_id, (SELECT COUNT(DISTINCT user_id) FROM attempts a WHERE a.test_id=t.id "
        "AND a.status='done') AS n FROM tests t ORDER BY t.group_id, t.sort")
    btns = [[(f"[{(await db.group_name(t['group_id'])).split()[0]}] {t['title']} ({t['n']})", f"adm:rs:{t['id']}")]
            for t in tests]
    btns.append([(f"📈 Tahlil: {g['name']}", f"adm:anl:{g['id']}") for g in await db.study_groups()])
    btns.append([("📥 Hammasini Excel'da yuklab olish", "adm:xlsx")])
    btns.append(BACK)
    await safe_edit(cb, "📊 <b>Natijalar</b>\nTestni tanlang (qavsda — ishlaganlar soni):", kb(btns))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:rs:"))
async def results_test(cb: CallbackQuery):
    tid = cb.data.split(":", 2)[2]
    t = await db.fetchone("SELECT * FROM tests WHERE id=?", (tid,))
    mode = await db.get_setting("ranking_mode")
    ranked = await db.ranking(tid, mode)
    lines = [f"📊 <b>{html.escape(t['title'])}</b>", f"Ishtirokchilar: {len(ranked)} · "
             + ("birinchi urinish bo'yicha" if mode == "first" else "eng yaxshi urinish bo'yicha"), ""]
    if ranked:
        avg = sum(pct(r["correct"], r["total"]) for r in ranked) / len(ranked)
        lines.append(f"O'rtacha natija: <b>{avg:.0f}%</b>\n")
    for i, r in enumerate(ranked, 1):
        lines.append(f"{i}. {html.escape(r['full_name'])} — <b>{r['correct']}/{r['total']}</b> "
                     f"({pct(r['correct'], r['total'])}%) · ⏱ {fmt_dur(r['duration'])}")
    if not ranked:
        lines.append("Hali hech kim ishlamagan.")
    parts = chunk_text("\n".join(lines))
    await safe_edit(cb, parts[0], kb([[("⬅️ Natijalar", "adm:res")]]) if len(parts) == 1 else None)
    for p in parts[1:]:
        await cb.message.answer(p)
    if len(parts) > 1:
        await cb.message.answer("…", reply_markup=kb([[("⬅️ Natijalar", "adm:res")]]))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:anl"))
async def analysis(cb: CallbackQuery):
    """Qaysi testlar kam ishlanmoqda va qaysi savollarda ko'p xato qilinmoqda (guruh bo'yicha)."""
    parts = cb.data.split(":")
    gid = int(parts[2]) if len(parts) > 2 else 1
    students = (await db.fetchone("SELECT COUNT(*) AS n FROM users WHERE status='approved' AND group_id=?",
                                  (gid,)))["n"]
    tests = await db.fetchall("SELECT * FROM tests WHERE is_open=1 AND group_id=? ORDER BY sort", (gid,))
    daily_today = (await db.fetchone("SELECT COUNT(*) AS n FROM daily_log d JOIN users u ON u.id=d.user_id "
                                     "WHERE d.day=? AND u.group_id=?",
                                     (datetime.now(TZ).date().isoformat(), gid)))["n"]
    lines = [f"📈 <b>Tahlil: {html.escape(await db.group_name(gid))}</b> · faol o'quvchilar: {students}",
             f"📅 Bugungi vazifani bajarganlar: {daily_today}/{students}\n",
             "<b>Testlar</b> (ishlaganlar · o'rtacha natija):"]
    for t in tests:
        r = await db.fetchone("SELECT COUNT(DISTINCT user_id) AS u, AVG(CAST(correct AS REAL)/total) AS a "
                              "FROM attempts WHERE test_id=? AND status='done'", (t["id"],))
        g = await db.fetchone("SELECT COUNT(*) AS n, SUM(ok) AS c FROM group_answers ga "
                              "JOIN questions q ON q.id=ga.question_id WHERE q.test_id=?", (t["id"],))
        share = round(100 * r["u"] / students) if students else 0
        flag = "🔴" if share < 50 else ("🟡" if share < 80 else "🟢")
        avg = f"{round(100 * r['a'])}%" if r["a"] is not None else "—"
        grp = f" · guruhda {round(100 * (g['c'] or 0) / g['n'])}% to'g'ri" if g["n"] else ""
        lines.append(f"{flag} {html.escape(t['title'])}: {r['u']}/{students} ({share}%) · {avg}{grp}")
    hard = await db.fetchall("""
        SELECT q.q, t.title, SUM(n) AS n, SUM(c) AS c FROM (
            SELECT question_id, COUNT(*) AS n, SUM(is_correct) AS c FROM answers GROUP BY question_id
            UNION ALL
            SELECT question_id, COUNT(*) AS n, SUM(ok) AS c FROM group_answers
            WHERE question_id IS NOT NULL GROUP BY question_id
        ) s JOIN questions q ON q.id = s.question_id JOIN tests t ON t.id = q.test_id
        WHERE t.group_id = ?
        GROUP BY s.question_id, q.q, t.title HAVING SUM(n) >= 3
        ORDER BY 1.0 * SUM(c) / SUM(n) ASC LIMIT 10""", (gid,))
    lines.append("\n<b>Eng ko'p xato qilinayotgan savollar:</b>")
    if hard:
        for i, h in enumerate(hard, 1):
            err = round(100 * (h["n"] - (h["c"] or 0)) / h["n"])
            lines.append(f"{i}. ❌ {err}% xato ({h['n']} javob) — {html.escape(h['q'][:90])}")
    else:
        lines.append("Hali yetarli ma'lumot yo'q (har bir savolga kamida 3 ta javob kerak).")
    lines.append("\n<i>🔴 o'quvchilarning yarmidan kami ishlagan · 🟡 80% gacha · 🟢 80% dan ko'p.\n"
                 "Guruhga avtomatik savollar aynan shu tahlil asosida tanlanadi: kam ishlangan testlar va "
                 "ko'p xato qilinayotgan savollar tez-tez chiqadi.</i>")
    parts = chunk_text("\n".join(lines))
    await safe_edit(cb, parts[0], kb([[("⬅️ Natijalar", "adm:res")]]) if len(parts) == 1 else None)
    for p in parts[1:]:
        await cb.message.answer(p, reply_markup=kb([[("⬅️ Natijalar", "adm:res")]]) if p is parts[-1] else None)
    await cb.answer()


def _autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, 50)


def _header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
    ws.freeze_panes = "A2"


@router.callback_query(F.data == "adm:xlsx")
async def export_xlsx(cb: CallbackQuery):
    await cb.answer("Tayyorlanmoqda…")
    wb = Workbook()
    ws = wb.active
    ws.title = "Barcha urinishlar"
    _header(ws, ["O'quvchi", "Guruh", "Test", "To'g'ri", "Jami", "Foiz", "Vaqt", "Sana", "Holat"])
    rows = await db.fetchall("SELECT a.*, u.full_name, u.group_id, t.title FROM attempts a JOIN users u ON u.id=a.user_id "
                             "JOIN tests t ON t.id=a.test_id WHERE a.status IN ('done','stopped') "
                             "ORDER BY a.finished_at")
    for r in rows:
        ws.append([r["full_name"], await db.group_name(r["group_id"]), r["title"], r["correct"], r["total"], pct(r["correct"], r["total"]),
                   fmt_dur(r["duration"]), datetime.fromtimestamp(r["finished_at"], TZ).strftime("%d.%m.%Y %H:%M"),
                   "tugallangan" if r["status"] == "done" else "to'xtatilgan"])
    _autosize(ws)
    mode = await db.get_setting("ranking_mode")
    used = set()
    for t in await db.fetchall("SELECT * FROM tests ORDER BY group_id, sort"):
        ranked = await db.ranking(t["id"], mode)
        if not ranked:
            continue
        prefix = (await db.group_name(t["group_id"])).split()[0] + "-"
        name = (prefix + re.sub(r"[\[\]:*?/\\]", "", t["title"]))[:28]
        while name in used:
            name = name[:26] + str(len(used))
        used.add(name)
        s = wb.create_sheet(name)
        _header(s, ["O'rin", "O'quvchi", "To'g'ri", "Jami", "Foiz", "Vaqt"])
        for i, r in enumerate(ranked, 1):
            s.append([i, r["full_name"], r["correct"], r["total"], pct(r["correct"], r["total"]), fmt_dur(r["duration"])])
        _autosize(s)
    us = wb.create_sheet("O'quvchilar")
    _header(us, ["Ism familiya", "Guruh", "Username", "Telegram ID", "Holat", "Ro'yxatdan o'tgan"])
    for u in await db.fetchall("SELECT * FROM users ORDER BY group_id, full_name"):
        us.append([u["full_name"], await db.group_name(u["group_id"]), u["username"] or "", u["id"], u["status"],
                   datetime.fromtimestamp(u["created_at"], TZ).strftime("%d.%m.%Y")])
    _autosize(us)
    buf = io.BytesIO()
    wb.save(buf)
    fname = f"natijalar_{datetime.now(TZ).strftime('%Y-%m-%d')}.xlsx"
    await cb.message.answer_document(BufferedInputFile(buf.getvalue(), fname), caption="📥 Barcha natijalar")


# ======================= Darslar =======================
@router.callback_query(F.data.startswith("adm:les"))
async def lessons_admin(cb: CallbackQuery, state: FSMContext | None = None, gid: int | None = None):
    if state:
        await state.clear()
    parts = cb.data.split(":")
    if gid is None:
        gid = int(parts[2]) if len(parts) > 2 and parts[1] == "les" and parts[2].isdigit() else 1
    rows = await db.fetchall("SELECT l.*, (SELECT COUNT(*) FROM lesson_pages p WHERE p.lesson_id=l.id) AS np "
                             "FROM lessons l WHERE group_id=? ORDER BY num", (gid,))
    btns = await group_buttons("adm:les", gid)
    btns += [[(f"{'🟢' if l['is_open'] else '🔴'} {l['num']}-dars: {l['short']}", f"adm:le:{l['id']}")] for l in rows]
    btns.append([("➕ Yangi dars qo'shish", f"adm:lenew:{gid}")])
    btns.append(BACK)
    await safe_edit(cb, f"📚 <b>Darslar: {html.escape(await db.group_name(gid))}</b>\n"
                        "🟢 ochiq · 🔴 yopiq (o'quvchilar ko'rmaydi)", kb(btns))
    await cb.answer()


async def lesson_card(cb_or_msg, lid: int):
    l = await db.fetchone("SELECT * FROM lessons WHERE id=?", (lid,))
    np = (await db.fetchone("SELECT COUNT(*) AS n FROM lesson_pages WHERE lesson_id=?", (lid,)))["n"]
    tests = await db.fetchall("SELECT title FROM tests WHERE lesson_id=?", (lid,))
    reads = (await db.fetchone("SELECT COUNT(*) AS n FROM lesson_reads WHERE lesson_id=?", (lid,)))["n"]
    text = (f"📖 <b>{l['num']}-dars: {html.escape(l['title'])}</b>\n\n"
            f"Guruh: <b>{html.escape(await db.group_name(l['group_id']))}</b>\n"
            f"Holati: {'🟢 ochiq' if l['is_open'] else '🔴 yopiq'}\n"
            f"Matn sahifalari: {np}\n"
            f"Biriktirilgan fayl: {html.escape(l['file_name']) if l['file_id'] else '—'}\n"
            f"Testlar: {', '.join(html.escape(t['title']) for t in tests) or '—'}\n"
            f"O'qiganlar: {reads} nafar")
    btns = kb([
        [("👁 Ko'rish", f"les:open:{lid}:0"), ("🔴 Yopish" if l["is_open"] else "🟢 Ochish", f"adm:letog:{lid}")],
        [("✏️ Matn/fayl yuklash", f"adm:leup:{lid}"), ("👥 Guruhini almashtirish", f"adm:legrp:{lid}")],
        [("🗑 Faylni olib tashlash", f"adm:lefdel:{lid}")] if l["file_id"] else [],
        [("🗑 Darsni o'chirish", f"adm:ledel:{lid}")],
        [("⬅️ Darslar", f"adm:les:{l['group_id']}")],
    ])
    if isinstance(cb_or_msg, CallbackQuery):
        await safe_edit(cb_or_msg, text, btns)
    else:
        await cb_or_msg.answer(text, reply_markup=btns)


@router.callback_query(F.data.startswith("adm:le:"))
async def lesson_admin(cb: CallbackQuery):
    await lesson_card(cb, int(cb.data.split(":")[2]))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:legrp:"))
async def lesson_group(cb: CallbackQuery):
    """Dars (va unga bog'langan testlar)ni keyingi guruhga o'tkazadi."""
    lid = int(cb.data.split(":")[2])
    l = await db.fetchone("SELECT group_id FROM lessons WHERE id=?", (lid,))
    ids = [g["id"] for g in await db.study_groups()]
    new = ids[(ids.index(l["group_id"]) + 1) % len(ids)] if l["group_id"] in ids else ids[0]
    await db.execute("UPDATE lessons SET group_id=? WHERE id=?", (new, lid))
    await db.execute("UPDATE tests SET group_id=? WHERE lesson_id=?", (new, lid))
    await lesson_card(cb, lid)
    await cb.answer(f"Endi: {await db.group_name(new)}")


@router.callback_query(F.data.startswith("adm:letog:"))
async def lesson_toggle(cb: CallbackQuery):
    lid = int(cb.data.split(":")[2])
    await db.execute("UPDATE lessons SET is_open=1-is_open WHERE id=?", (lid,))
    await lesson_card(cb, lid)
    await cb.answer("Saqlandi")


@router.callback_query(F.data.startswith("adm:lefdel:"))
async def lesson_file_del(cb: CallbackQuery):
    lid = int(cb.data.split(":")[2])
    await db.execute("UPDATE lessons SET file_id=NULL, file_name=NULL WHERE id=?", (lid,))
    await lesson_card(cb, lid)
    await cb.answer("Fayl olib tashlandi")


@router.callback_query(F.data.startswith("adm:ledel:"))
async def lesson_del(cb: CallbackQuery):
    parts = cb.data.split(":")
    lid = int(parts[2])
    if len(parts) < 4:
        await safe_edit(cb, "❗ Darsni o'chirasizmi? Matni o'chadi, unga bog'langan testlar esa saqlanib qoladi.",
                        kb([[("🗑 Ha, o'chirish", f"adm:ledel:{lid}:yes"), ("Bekor qilish", f"adm:le:{lid}")]]))
        return await cb.answer()
    await db.execute("DELETE FROM lesson_pages WHERE lesson_id=?", (lid,))
    await db.execute("DELETE FROM lesson_reads WHERE lesson_id=?", (lid,))
    await db.execute("UPDATE tests SET lesson_id=NULL WHERE lesson_id=?", (lid,))
    l = await db.fetchone("SELECT group_id FROM lessons WHERE id=?", (lid,))
    await db.execute("DELETE FROM lessons WHERE id=?", (lid,))
    await cb.answer("O'chirildi")
    await lessons_admin(cb, gid=l["group_id"] if l else 1)


@router.callback_query(F.data.startswith("adm:lenew"))
async def lesson_new(cb: CallbackQuery, state: FSMContext):
    parts = cb.data.split(":")
    await state.set_state(AdmSt.lesson_title)
    await state.update_data(gid=int(parts[2]) if len(parts) > 2 else 1)
    await cb.message.answer("✍️ Yangi dars nomini yozing (masalan: <i>Kelasi zamon</i>).\nBekor qilish: /cancel")
    await cb.answer()


@router.message(AdmSt.lesson_title, F.text)
async def lesson_new_title(msg: Message, state: FSMContext):
    if await _left_input(msg, state):
        return
    title = msg.text.strip()[:120]
    gid = (await state.get_data()).get("gid", 1)
    num = ((await db.fetchone("SELECT MAX(num) AS m FROM lessons WHERE group_id=?", (gid,)))["m"] or 0) + 1
    lid = await db.execute("INSERT INTO lessons(num, title, short, group_id) VALUES(?,?,?,?)",
                           (num, title, title[:40], gid))
    await state.set_state(AdmSt.lesson_content)
    await state.update_data(lid=lid, replace=True)
    await msg.answer(f"✅ {num}-dars yaratildi: <b>{html.escape(title)}</b>\n\n" + LESSON_UPLOAD_HELP)


LESSON_UPLOAD_HELP = ("Endi darslik matnini yuboring:\n"
                      "• <b>Word (.docx)</b> fayl — matni avtomatik sahifalarga bo'linadi va fayl ham biriktiriladi;\n"
                      "• <b>PDF yoki boshqa fayl</b> — o'quvchilar yuklab oladigan fayl sifatida biriktiriladi;\n"
                      "• <b>oddiy matn</b> — bir nechta xabar yuborsangiz, har biri alohida sahifa bo'ladi.\n\n"
                      "Tugatgach «✅ Tayyor» tugmasini bosing. Bekor qilish: /cancel")


@router.callback_query(F.data.startswith("adm:leup:"))
async def lesson_upload(cb: CallbackQuery, state: FSMContext):
    lid = int(cb.data.split(":")[2])
    await state.set_state(AdmSt.lesson_content)
    await state.update_data(lid=lid, replace=True)
    await cb.message.answer("✏️ Yangi matn yuborsangiz, eski matn <b>almashtiriladi</b>.\n\n" + LESSON_UPLOAD_HELP,
                            reply_markup=kb([[("✅ Tayyor", f"adm:ledone:{lid}")]]))
    await cb.answer()


async def _lesson_add_pages(state: FSMContext, lid: int, pages: list[str]):
    data = await state.get_data()
    if data.get("replace"):
        await db.execute("DELETE FROM lesson_pages WHERE lesson_id=?", (lid,))
        await state.update_data(replace=False)
    start = (await db.fetchone("SELECT COALESCE(MAX(idx), -1) AS m FROM lesson_pages WHERE lesson_id=?", (lid,)))["m"] + 1
    for i, p in enumerate(pages):
        await db.execute("INSERT INTO lesson_pages(lesson_id, idx, html) VALUES(?,?,?)", (lid, start + i, p))


@router.message(AdmSt.lesson_content, F.document)
async def lesson_upload_doc(msg: Message, state: FSMContext, bot: Bot):
    lid = (await state.get_data())["lid"]
    doc = msg.document
    name = doc.file_name or "fayl"
    await db.execute("UPDATE lessons SET file_id=?, file_name=? WHERE id=?", (doc.file_id, name, lid))
    note = "📎 Fayl darsga biriktirildi."
    if name.lower().endswith(".docx"):
        f = await bot.download(doc)
        try:
            pages = paginate(lesson_blocks(f.read()))
            await _lesson_add_pages(state, lid, pages)
            note += f"\n📄 Matn olindi: {len(pages)} sahifa."
        except Exception as e:  # noqa: BLE001
            note += f"\n⚠️ Matnni o'qib bo'lmadi: {e}"
    await msg.answer(note + "\nYana yuborishingiz yoki «✅ Tayyor» ni bosishingiz mumkin.",
                     reply_markup=kb([[("✅ Tayyor", f"adm:ledone:{lid}")]]))


@router.message(AdmSt.lesson_content, F.text)
async def lesson_upload_text(msg: Message, state: FSMContext):
    if await _left_input(msg, state):
        return
    lid = (await state.get_data())["lid"]
    await _lesson_add_pages(state, lid, paginate([msg.html_text]))
    await msg.answer("📄 Sahifa qo'shildi. Yana yuboring yoki «✅ Tayyor» ni bosing.",
                     reply_markup=kb([[("✅ Tayyor", f"adm:ledone:{lid}")]]))


@router.callback_query(F.data.startswith("adm:ledone:"))
async def lesson_upload_done(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    await lesson_card(cb.message, int(cb.data.split(":")[2]))
    await cb.answer("Saqlandi")


# ======================= Testlar =======================
@router.callback_query(F.data.startswith("adm:tst"))
async def tests_admin(cb: CallbackQuery, state: FSMContext | None = None, gid: int | None = None):
    if state:
        await state.clear()
    parts = cb.data.split(":")
    if gid is None:
        gid = int(parts[2]) if len(parts) > 2 and parts[1] == "tst" and parts[2].isdigit() else 1
    rows = await db.fetchall("SELECT t.*, (SELECT COUNT(*) FROM questions q WHERE q.test_id=t.id) AS nq "
                             "FROM tests t WHERE group_id=? ORDER BY sort, title", (gid,))
    btns = await group_buttons("adm:tst", gid)
    btns += [[(f"{'🟢' if t['is_open'] else '🔴'} {t['title']} [{t['nq']}]", f"adm:te:{t['id']}")] for t in rows]
    btns.append([("➕ Word'dan yangi test yuklash", f"adm:tenew:{gid}")])
    btns.append(BACK)
    await safe_edit(cb, f"📝 <b>Testlar: {html.escape(await db.group_name(gid))}</b>\n"
                        "🟢 ochiq · 🔴 yopiq · [savollar soni]", kb(btns))
    await cb.answer()


async def test_card(target, tid: str):
    t = await db.fetchone("SELECT t.*, l.num AS lnum, l.short AS lshort FROM tests t "
                          "LEFT JOIN lessons l ON l.id=t.lesson_id WHERE t.id=?", (tid,))
    nq = (await db.fetchone("SELECT COUNT(*) AS n FROM questions WHERE test_id=?", (tid,)))["n"]
    per = t["per_attempt"]
    text = (f"📝 <b>{html.escape(t['title'])}</b>\n\n"
            f"Guruh: <b>{html.escape(await db.group_name(t['group_id']))}</b>\n"
            f"Holati: {'🟢 ochiq' if t['is_open'] else '🔴 yopiq'}\n"
            f"Savollar bazasi: {nq} ta\n"
            f"Bitta urinishda: {per if per else 'hammasi'} ta (tasodifiy tanlanadi)\n"
            f"Dars: {str(t['lnum']) + '-dars (' + html.escape(t['lshort']) + ')' if t['lnum'] else 'bog‘lanmagan'}\n"
            f"Manba: {html.escape(t['source'] or '—')}")
    per_row = [(("🔘 " if per == n else "") + (str(n) if n else "Hammasi"), f"adm:teper:{tid}:{n}")
               for n in PER_ATTEMPT_CHOICES]
    btns = kb([
        [("🔴 Yopish" if t["is_open"] else "🟢 Ochish", f"adm:tetog:{tid}"), ("📄 Word'da yuklab olish", f"adm:tedoc:{tid}")],
        [("Bitta urinishdagi savollar soni:", "noop")], per_row,
        [("🔗 Darsga bog'lash", f"adm:telink:{tid}"), ("♻️ Word bilan almashtirish", f"adm:terep:{tid}")],
        [("👥 Guruhini almashtirish", f"adm:tegrp:{tid}"), ("🗑 Testni o'chirish", f"adm:tedel:{tid}")],
        [("⬅️ Testlar", f"adm:tst:{t['group_id']}")],
    ])
    if isinstance(target, CallbackQuery):
        await safe_edit(target, text, btns)
    else:
        await target.answer(text, reply_markup=btns)


@router.callback_query(F.data.startswith("adm:te:"))
async def test_admin(cb: CallbackQuery):
    await test_card(cb, cb.data.split(":", 2)[2])
    await cb.answer()


@router.callback_query(F.data.startswith("adm:tegrp:"))
async def test_group(cb: CallbackQuery):
    tid = cb.data.split(":", 2)[2]
    t = await db.fetchone("SELECT group_id FROM tests WHERE id=?", (tid,))
    ids = [g["id"] for g in await db.study_groups()]
    new = ids[(ids.index(t["group_id"]) + 1) % len(ids)] if t["group_id"] in ids else ids[0]
    await db.execute("UPDATE tests SET group_id=?, lesson_id=NULL WHERE id=?", (new, tid))
    await test_card(cb, tid)
    await cb.answer(f"Endi: {await db.group_name(new)} (dars bog'lanishi olib tashlandi)", show_alert=True)


@router.callback_query(F.data.startswith("adm:tetog:"))
async def test_toggle(cb: CallbackQuery):
    tid = cb.data.split(":", 2)[2]
    await db.execute("UPDATE tests SET is_open=1-is_open WHERE id=?", (tid,))
    await test_card(cb, tid)
    await cb.answer("Saqlandi")


@router.callback_query(F.data.startswith("adm:teper:"))
async def test_per(cb: CallbackQuery):
    _, _, tid, n = cb.data.split(":")
    await db.execute("UPDATE tests SET per_attempt=? WHERE id=?", (int(n), tid))
    await test_card(cb, tid)
    await cb.answer("Saqlandi")


@router.callback_query(F.data.startswith("adm:tedoc:"))
async def test_doc(cb: CallbackQuery):
    tid = cb.data.split(":", 2)[2]
    t = await db.fetchone("SELECT * FROM tests WHERE id=?", (tid,))
    qs = await db.test_questions(tid)
    data = export_tests(t["title"], qs, "To'g'ri javoblar qalin (bold) qilib belgilangan. Faylni tahrirlab, "
                                        "botga «♻️ Word bilan almashtirish» orqali qayta yuklashingiz mumkin.")
    safe = re.sub(r"[^\w\-]+", "_", t["title"])[:50]
    await cb.message.answer_document(BufferedInputFile(data, f"{safe}.docx"), caption=f"📄 {t['title']} — {len(qs)} ta savol")
    await cb.answer()


@router.callback_query(F.data.startswith("adm:telink:"))
async def test_link(cb: CallbackQuery):
    parts = cb.data.split(":")
    tid = parts[2]
    if len(parts) == 4:
        lid = int(parts[3]) or None
        await db.execute("UPDATE tests SET lesson_id=? WHERE id=?", (lid, tid))
        await test_card(cb, tid)
        return await cb.answer("Saqlandi")
    t = await db.fetchone("SELECT group_id FROM tests WHERE id=?", (tid,))
    lessons = await db.fetchall("SELECT * FROM lessons WHERE group_id=? ORDER BY num", (t["group_id"],))
    btns = [[(f"{l['num']}-dars: {l['short']}", f"adm:telink:{tid}:{l['id']}")] for l in lessons]
    btns.append([("Hech qaysi darsga (umumiy test)", f"adm:telink:{tid}:0")])
    btns.append([("⬅️ Orqaga", f"adm:te:{tid}")])
    await safe_edit(cb, "🔗 Test qaysi darsdan keyin ochilsin?", kb(btns))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:tedel:"))
async def test_del(cb: CallbackQuery):
    parts = cb.data.split(":")
    tid = parts[2]
    if len(parts) < 4:
        await safe_edit(cb, "❗ Testni o'chirasizmi? Savollar va shu testning barcha natijalari o'chadi.",
                        kb([[("🗑 Ha, o'chirish", f"adm:tedel:{tid}:yes"), ("Bekor qilish", f"adm:te:{tid}")]]))
        return await cb.answer()
    await db.execute("DELETE FROM answers WHERE attempt_id IN (SELECT id FROM attempts WHERE test_id=?)", (tid,))
    await db.execute("DELETE FROM attempts WHERE test_id=?", (tid,))
    t = await db.fetchone("SELECT group_id FROM tests WHERE id=?", (tid,))
    await db.execute("DELETE FROM questions WHERE test_id=?", (tid,))
    await db.execute("DELETE FROM tests WHERE id=?", (tid,))
    await cb.answer("O'chirildi")
    await tests_admin(cb, gid=t["group_id"] if t else 1)


IMPORT_HELP = ("📄 Word (.docx) faylni yuboring.\n\n<b>Fayl formati:</b>\n"
               "<code>1. Savol matni?\nA) variant\nB) to'g'ri variant\nC) variant\nD) variant</code>\n\n"
               "To'g'ri javobni <b>qalin (bold)</b> qilib belgilang. Yoki fayl oxiriga kalit yozing: "
               "<code>1 — B, 2 — A ...</code>\nIxtiyoriy izoh: variantlardan keyin <code>Izoh: ...</code> qatori.\n\n"
               "Bekor qilish: /cancel")


@router.callback_query(F.data.startswith("adm:tenew"))
async def test_new(cb: CallbackQuery, state: FSMContext):
    parts = cb.data.split(":")
    await state.set_state(AdmSt.test_file)
    await state.update_data(gid=int(parts[2]) if len(parts) > 2 else 1)
    await cb.message.answer(IMPORT_HELP)
    await cb.answer()


@router.callback_query(F.data.startswith("adm:terep:"))
async def test_replace(cb: CallbackQuery, state: FSMContext):
    tid = cb.data.split(":", 2)[2]
    await state.set_state(AdmSt.test_replace)
    await state.update_data(tid=tid)
    await cb.message.answer("♻️ Yangi Word faylni yuboring — testning savollari shu fayldagilar bilan "
                            "<b>almashtiriladi</b> (eski natijalar saqlanadi).\n\n" + IMPORT_HELP)
    await cb.answer()


async def _read_test_file(msg: Message, bot: Bot):
    doc = msg.document
    if not (doc.file_name or "").lower().endswith(".docx"):
        await msg.answer("❗ Faqat Word (.docx) fayl qabul qilinadi. Eski .doc bo'lsa, Word'da «Save as → .docx» qiling.")
        return None
    f = await bot.download(doc)
    try:
        qs, issues = parse_tests(f.read())
    except Exception as e:  # noqa: BLE001
        await msg.answer(f"❗ Faylni o'qib bo'lmadi: {e}")
        return None
    report = f"🔎 Topilgan savollar: <b>{len(qs)}</b> ta"
    if issues:
        report += f"\n⚠️ Qabul qilinmadi: {len(issues)} ta\n" + "\n".join("• " + html.escape(i) for i in issues[:15])
        if len(issues) > 15:
            report += f"\n… va yana {len(issues) - 15} ta"
    await msg.answer(report)
    if not qs:
        await msg.answer("Savol topilmadi. Formatni tekshirib, qayta yuboring.\n\n" + IMPORT_HELP)
        return None
    return qs


@router.message(AdmSt.test_file, F.document)
async def test_new_file(msg: Message, state: FSMContext, bot: Bot):
    qs = await _read_test_file(msg, bot)
    if qs is None:
        return
    await state.update_data(qs=qs, fname=msg.document.file_name)
    await state.set_state(AdmSt.test_title)
    await msg.answer("✍️ Test nomini yozing (masalan: <i>8-dars testi: Kelasi zamon</i>):")


@router.message(AdmSt.test_title, F.text)
async def test_new_title(msg: Message, state: FSMContext):
    if await _left_input(msg, state):
        return
    data = await state.get_data()
    title = msg.text.strip()[:120]
    tid = f"u{int(time.time())}"
    sort = ((await db.fetchone("SELECT MAX(sort) AS m FROM tests"))["m"] or 0) + 10
    await db.execute("INSERT INTO tests(id, title, sort, per_attempt, source, group_id) VALUES(?,?,?,?,?,?)",
                     (tid, title, sort, 30, f"Word fayl: {data.get('fname', '')}", data.get("gid", 1)))
    await db.replace_questions(tid, data["qs"])
    await state.clear()
    await msg.answer(f"✅ Test qo'shildi: <b>{html.escape(title)}</b> — {len(data['qs'])} ta savol.\n"
                     "Endi uni darsga bog'lang (dars o'qilgach ochiladi) yoki umumiy test sifatida qoldiring.")
    await test_card(msg, tid)


@router.message(AdmSt.test_replace, F.document)
async def test_replace_file(msg: Message, state: FSMContext, bot: Bot):
    qs = await _read_test_file(msg, bot)
    if qs is None:
        return
    tid = (await state.get_data())["tid"]
    await db.replace_questions(tid, qs)
    await state.clear()
    await msg.answer(f"✅ Savollar almashtirildi: {len(qs)} ta.")
    await test_card(msg, tid)


@router.message(StateFilter(AdmSt.test_file, AdmSt.test_replace))
async def test_expect_file(msg: Message):
    await msg.answer("Word (.docx) fayl yuboring yoki /cancel bosing.")


# ======================= Guruh sozlamalari =======================
HOUR_PRESETS = [(7, 21), (8, 21), (8, 22), (9, 20), (0, 24)]


@router.callback_query(F.data == "adm:grp")
async def group_menu(cb: CallbackQuery):
    kick_scheduler()
    s = {k: await db.get_setting(k) for k in
         ("group_enabled", "group_interval", "group_start_hour", "group_end_hour", "group_fun", "group_fun_every",
          "group_scope", "group_seconds")}
    groups = await db.fetchall("SELECT * FROM groups ORDER BY added_at")
    scope = "🧠 aqlli tanlash (barcha ochiq testlar)"
    if s["group_scope"] != "all":
        t = await db.fetchone("SELECT title FROM tests WHERE id=?", (s["group_scope"],))
        scope = t["title"] if t else scope
    iv = int(s["group_interval"])
    chat_lines = [f"{'🟢' if g['active'] else '🔴'} {html.escape(g['title'] or str(g['chat_id']))} · "
                  f"{html.escape(await db.group_name(g['study_group_id'] or 1))} · ⏰ {g['sched_times'] or 'jadval yoq'}"
                  for g in groups]
    text = ("👥 <b>Guruh sozlamalari</b>\n\n"
            "<b>1) Jadval bo'yicha test</b> — har bir guruh uchun alohida: guruh nomini bosing va vaqtni "
            "belgilang (masalan, 16:00). Bot testni o'zi tanlaydi: avval yuborilmaganlarini dars tartibida, "
            "keyin eng ko'p xato qilinganlarini.\n\n"
            "<b>2) Oraliq bilan bitta-bitta savol</b> (barcha guruhlar uchun):\n"
            f"Holati: {yn(s['group_enabled'])} {'yoqilgan' if int(s['group_enabled']) else 'o‘chirilgan'}\n"
            f"Har <b>{iv if iv < 60 else iv // 60} {'daqiqada' if iv < 60 else 'soatda'}</b> bitta savol\n"
            f"Vaqt oralig'i: <b>{s['group_start_hour']}:00 – {s['group_end_hour']}:00</b>\n"
            f"Savolga vaqt: <b>{s['group_seconds']} soniya</b>\n"
            f"Qiziqarli savollar: {yn(s['group_fun'])} (har {s['group_fun_every']}-savol)\n"
            f"Savollar manbasi: <b>{html.escape(scope)}</b>\n\n"
            f"Guruhlar ({len(groups)}):\n" +
            ("\n".join(chat_lines) or "— hali yo'q. Botni guruhga qo'shing va guruhda /savol yozing.") +
            "\n\n<i>Guruhda: /reyting — reyting, /test — test tanlab yuborish, /savol — bitta savol, /stoptest — to'xtatish.</i>")
    btns = [[(("⏸ O'chirish" if int(s["group_enabled"]) else "▶️ Yoqish"), "adm:gen")],
            [(("🔘 " if iv == m else "") + (f"{m}d" if m < 60 else f"{m // 60}s"), f"adm:gint:{m}")
             for m in INTERVAL_CHOICES],
            [(("🔘 " if (str(a), str(b)) == (s["group_start_hour"], s["group_end_hour"]) else "") + f"{a}–{b}",
              f"adm:ghr:{a}:{b}") for a, b in HOUR_PRESETS],
            [(f"🧠 Qiziqarli: {yn(s['group_fun'])}", "adm:gfun"), ("📚 Manba", "adm:gscope")],
            [("📝 Guruhga test yuborish", "adm:gt"), ("📤 Bitta savol", "adm:gnow")]]
    for chat_id in list(running):
        btns.append([(f"⏹ Ketayotgan testni to'xtatish ({running[chat_id].title[:18]})", f"adm:gtstop:{chat_id}")])
    for g in groups:
        title = (g["title"] or str(g["chat_id"]))[:28]
        btns.append([(f"{'🟢' if g['active'] else '🔴'} {title} ⚙️", f"adm:gch:{g['chat_id']}")])
    btns.append(BACK)
    await safe_edit(cb, text, kb(btns))
    await cb.answer()


async def chat_card(cb: CallbackQuery, chat_id: int):
    g = await db.fetchone("SELECT * FROM groups WHERE chat_id=?", (chat_id,))
    if not g:
        return await cb.answer("Guruh topilmadi", show_alert=True)
    times = g["sched_times"] or ""
    cnt = g["sched_count"] or 10
    last = await db.fetchone("SELECT gs.sent_at, t.title FROM group_sessions gs LEFT JOIN tests t ON t.id=gs.test_id "
                             "WHERE gs.chat_id=? ORDER BY gs.sent_at DESC LIMIT 1", (chat_id,))
    sent_n = (await db.fetchone("SELECT COUNT(DISTINCT test_id) AS n FROM group_sessions WHERE chat_id=?",
                                (chat_id,)))["n"]
    total_n = (await db.fetchone("SELECT COUNT(*) AS n FROM tests WHERE is_open=1 AND group_id=?",
                                 (g["study_group_id"] or 1,)))["n"]
    text = (f"👥 <b>{html.escape(g['title'] or str(chat_id))}</b>\n\n"
            f"Holati: {'🟢 faol' if g['active'] else '🔴 o‘chirilgan'}\n"
            f"O'quv guruhi: <b>{html.escape(await db.group_name(g['study_group_id'] or 1))}</b> "
            "(shu guruh testlari yuboriladi)\n"
            f"⏰ Jadval: <b>{times or 'belgilanmagan'}</b> · har safar {cnt} ta savol\n"
            f"📊 Yuborilgan testlar: {sent_n} / {total_n}\n"
            f"Oxirgisi: {html.escape(last['title'] or '—') + ' · ' + datetime.fromtimestamp(last['sent_at'], TZ).strftime('%d.%m %H:%M') if last else '—'}")
    btns = [[("🔴 O'chirish" if g["active"] else "🟢 Yoqish", f"adm:gtog:{chat_id}"),
             ("👥 O'quv guruhini almashtirish", f"adm:gsg:{chat_id}")],
            [("⏰ Jadval vaqtini o'zgartirish", f"adm:gsch:{chat_id}")],
            [(("🔘 " if cnt == n else "") + f"{n} ta", f"adm:gscn:{chat_id}:{n}") for n in (5, 10, 15, 20, 30)],
            [("📝 Hozir test yuborish", f"adm:gt:{chat_id}"), ("🏆 Reyting", f"adm:gtop:{chat_id}")],
            [("♻️ Reytingni nolga tushirish", f"adm:grst:{chat_id}")],
            [("⬅️ Guruh sozlamalari", "adm:grp")]]
    await safe_edit(cb, text, kb(btns))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:gch:"))
async def chat_card_cb(cb: CallbackQuery):
    await chat_card(cb, int(cb.data.split(":")[2]))


@router.callback_query(F.data.startswith("adm:gsg:"))
async def chat_study_group(cb: CallbackQuery):
    chat_id = int(cb.data.split(":")[2])
    g = await db.fetchone("SELECT study_group_id FROM groups WHERE chat_id=?", (chat_id,))
    ids = [x["id"] for x in await db.study_groups()]
    cur = g["study_group_id"] or 1
    new = ids[(ids.index(cur) + 1) % len(ids)] if cur in ids else ids[0]
    await db.execute("UPDATE groups SET study_group_id=? WHERE chat_id=?", (new, chat_id))
    kick_scheduler()
    await chat_card(cb, chat_id)


@router.callback_query(F.data.startswith("adm:gscn:"))
async def chat_sched_count(cb: CallbackQuery):
    _, _, chat_id, n = cb.data.split(":")
    await db.execute("UPDATE groups SET sched_count=? WHERE chat_id=?", (int(n), int(chat_id)))
    kick_scheduler()
    await chat_card(cb, int(chat_id))


@router.callback_query(F.data.startswith("adm:gsch:"))
async def chat_sched_edit(cb: CallbackQuery, state: FSMContext):
    chat_id = int(cb.data.split(":")[2])
    await state.set_state(AdmSt.time_input)
    await state.update_data(key=f"sched:{chat_id}")
    await cb.message.answer("⏰ Guruhga test yuboriladigan vaqt(lar)ni yozing.\n"
                            "Masalan: <code>16:00</code> yoki bir nechta: <code>16:00, 19:30</code>\n"
                            "Jadvalni o'chirish uchun <code>0</code> yozing. Bekor qilish: /cancel")
    await cb.answer()


@router.message(AdmSt.time_input, F.text)
async def time_input(msg: Message, state: FSMContext):
    from group import parse_times
    if await _left_input(msg, state):
        return
    key = (await state.get_data()).get("key", "")
    times = parse_times(msg.text)
    if key.startswith("sched:"):
        chat_id = int(key.split(":")[1])
        if msg.text.strip() in ("0", "yo'q", "yoq", "-") :
            times = []
        elif not times:
            return await msg.answer("Vaqtni <code>16:00</code> ko'rinishida yozing yoki /cancel bosing.")
        await db.execute("UPDATE groups SET sched_times=? WHERE chat_id=?", (", ".join(times) or None, chat_id))
        kick_scheduler()
        await state.clear()
        return await msg.answer(f"✅ Jadval saqlandi: <b>{', '.join(times) or 'o‘chirildi'}</b>",
                                reply_markup=kb([[("⬅️ Guruh kartasi", f"adm:gch:{chat_id}")]]))
    if not times:
        return await msg.answer("Vaqtni <code>16:00</code> ko'rinishida yozing yoki /cancel bosing.")
    await db.set_setting(key, times[0])
    await state.clear()
    back = {"daily_time": "adm:dly", "remind_time": "adm:dly", "report_time": "adm:par"}.get(key, "adm:home")
    await msg.answer(f"✅ Saqlandi: <b>{times[0]}</b>", reply_markup=kb([[("⬅️ Orqaga", back)]]))


@router.callback_query(F.data.startswith("adm:gt:") | (F.data == "adm:gt"))
async def g_test(cb: CallbackQuery, bot: Bot):
    """Guruhga test yuborish: guruh → test → savollar soni."""
    parts = cb.data.split(":")
    groups = await db.fetchall("SELECT * FROM groups WHERE active=1 ORDER BY added_at")
    if not groups:
        return await cb.answer("Faol guruh yo'q. Botni guruhga qo'shing.", show_alert=True)
    if len(parts) == 2:
        if len(groups) == 1:
            parts = ["adm", "gt", str(groups[0]["chat_id"])]
        else:
            btns = [[(g["title"] or str(g["chat_id"]), f"adm:gt:{g['chat_id']}")] for g in groups]
            btns.append([("⬅️ Orqaga", "adm:grp")])
            await safe_edit(cb, "👥 Qaysi guruhga test yuboramiz?", kb(btns))
            return await cb.answer()
    chat_id = int(parts[2])
    if len(parts) == 3:
        from group import chat_group
        tests = await db.fetchall("SELECT t.id, t.title, (SELECT COUNT(*) FROM questions q WHERE q.test_id=t.id) AS n "
                                  "FROM tests t WHERE t.group_id=? ORDER BY t.sort", (await chat_group(chat_id),))
        btns = [[("🤖 Bot o'zi tanlasin (tahlil asosida)", f"adm:gt:{chat_id}:auto")]]
        btns += [[(f"{t['title']} [{t['n']}]", f"adm:gt:{chat_id}:{t['id']}")] for t in tests if t["n"]]
        btns.append([("⬅️ Orqaga", f"adm:gch:{chat_id}")])
        await safe_edit(cb, "📝 Guruhga qaysi test yuborilsin?", kb(btns))
        return await cb.answer()
    tid = parts[3]
    if tid == "auto":
        from group import pick_sched_test
        tid = await pick_sched_test(chat_id)
        if not tid:
            return await cb.answer("Bu guruh uchun ochiq test yo'q.", show_alert=True)
        parts[3] = tid
    if len(parts) == 4:
        t = await db.fetchone("SELECT title FROM tests WHERE id=?", (tid,))
        n = (await db.fetchone("SELECT COUNT(*) AS n FROM questions WHERE test_id=?", (tid,)))["n"]
        sec = await db.get_int("group_seconds")
        await safe_edit(cb, f"📝 <b>{html.escape(t['title'])}</b>\nNechta savol yuborilsin?\n"
                            f"<i>Har bir savolga {sec} soniya beriladi (⏱ Vaqt sozlamalaridan o'zgartiriladi).</i>",
                        kb(count_buttons(f"adm:gt:{chat_id}:{tid}", n) + [[("⬅️ Orqaga", f"adm:gt:{chat_id}")]]))
        return await cb.answer()
    err = await start_group_test(bot, chat_id, tid, int(parts[4]))
    if err:
        return await cb.answer(err, show_alert=True)
    await cb.answer("✅ Test guruhga yuborilmoqda!", show_alert=True)
    await group_menu(cb)


@router.callback_query(F.data.startswith("adm:gtstop:"))
async def g_test_stop(cb: CallbackQuery):
    stop_group_test(int(cb.data.split(":")[2]))
    await cb.answer("⏹ Test to'xtatiladi (joriy savol tugagach).", show_alert=True)


@router.callback_query(F.data == "adm:gen")
async def g_enable(cb: CallbackQuery):
    await db.set_setting("group_enabled", 1 - await db.get_int("group_enabled"))
    await group_menu(cb)


@router.callback_query(F.data.startswith("adm:gint:"))
async def g_interval(cb: CallbackQuery):
    await db.set_setting("group_interval", int(cb.data.split(":")[2]))
    await group_menu(cb)


@router.callback_query(F.data.startswith("adm:ghr:"))
async def g_hours(cb: CallbackQuery):
    _, _, a, b = cb.data.split(":")
    await db.set_setting("group_start_hour", a)
    await db.set_setting("group_end_hour", b)
    await group_menu(cb)


@router.callback_query(F.data == "adm:gfun")
async def g_fun(cb: CallbackQuery):
    await db.set_setting("group_fun", 1 - await db.get_int("group_fun"))
    await group_menu(cb)


@router.callback_query(F.data == "adm:gscope")
async def g_scope(cb: CallbackQuery):
    tests = await db.fetchall("SELECT * FROM tests ORDER BY sort")
    cur = await db.get_setting("group_scope")
    btns = [[(("🔘 " if cur == "all" else "") + "🧠 Aqlli tanlash (barcha ochiq testlar)", "adm:gsc:all")]]
    btns += [[(("🔘 " if cur == t["id"] else "") + t["title"], f"adm:gsc:{t['id']}")] for t in tests]
    btns.append([("⬅️ Orqaga", "adm:grp")])
    await safe_edit(cb, "📚 Guruhga qaysi test savollari tashlansin?", kb(btns))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:gsc:"))
async def g_scope_set(cb: CallbackQuery):
    await db.set_setting("group_scope", cb.data.split(":", 2)[2])
    await group_menu(cb)


@router.callback_query(F.data == "adm:gnow")
async def g_now(cb: CallbackQuery, bot: Bot):
    groups = await db.fetchall("SELECT * FROM groups WHERE active=1")
    if not groups:
        return await cb.answer("Faol guruh yo'q. Botni guruhga qo'shing.", show_alert=True)
    n = 0
    for g in groups:
        n += await post_question(bot, g["chat_id"])
    await cb.answer(f"📤 {n} ta guruhga savol yuborildi", show_alert=True)


@router.callback_query(F.data.startswith("adm:gtog:") | F.data.startswith("adm:gon:") | F.data.startswith("adm:goff:"))
async def g_toggle(cb: CallbackQuery):
    kind, chat_id = cb.data.split(":")[1], int(cb.data.split(":")[2])
    if kind == "gtog":
        await db.execute("UPDATE groups SET active=1-active WHERE chat_id=?", (chat_id,))
        kick_scheduler()
        return await chat_card(cb, chat_id)
    await db.execute("UPDATE groups SET active=? WHERE chat_id=?", (1 if kind == "gon" else 0, chat_id))
    kick_scheduler()
    await safe_edit(cb, cb.message.html_text + ("\n\n✅ Yoqildi" if kind == "gon" else "\n\n❌ O'chirildi"))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:gtop:"))
async def g_top(cb: CallbackQuery, bot: Bot):
    chat_id = int(cb.data.split(":")[2])
    text = await group_top_text(chat_id, 15)
    await cb.message.answer(text, reply_markup=kb([[("📣 Guruhga yuborish", f"adm:gpub:{chat_id}")]]))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:gpub:"))
async def g_pub(cb: CallbackQuery, bot: Bot):
    chat_id = int(cb.data.split(":")[2])
    try:
        await bot.send_message(chat_id, await group_top_text(chat_id, 10))
        await cb.answer("📣 Yuborildi")
    except Exception as e:  # noqa: BLE001
        await cb.answer(f"Yuborib bo'lmadi: {e}", show_alert=True)


@router.callback_query(F.data.startswith("adm:grst:"))
async def g_reset(cb: CallbackQuery):
    parts = cb.data.split(":")
    chat_id = int(parts[2])
    if len(parts) < 4:
        await safe_edit(cb, "♻️ Shu guruh reytingini nolga tushirasizmi?",
                        kb([[("Ha", f"adm:grst:{chat_id}:yes"), ("Yo'q", "adm:grp")]]))
        return await cb.answer()
    await db.execute("DELETE FROM group_scores WHERE chat_id=?", (chat_id,))
    await cb.answer("Reyting tozalandi")
    await group_menu(cb)


# ======================= Boshqa sozlamalar =======================
@router.callback_query(F.data == "adm:opt")
async def options(cb: CallbackQuery):
    rr, rm, rt = (await db.get_int("require_read"), await db.get_setting("ranking_mode"),
                  await db.get_int("allow_retake"))
    text = ("⚙️ <b>Boshqa sozlamalar</b>\n\n"
            f"📖 Testdan oldin darsni o'qish majburiy: {yn(rr)}\n"
            f"🔁 Testni qayta ishlashga ruxsat: {yn(rt)}\n"
            f"🏆 Reytingda hisoblanadi: <b>{'birinchi urinish' if rm == 'first' else 'eng yaxshi urinish'}</b>\n\n"
            "<i>Birinchi urinish — eng adolatli: bola javoblarni yodlab olib qayta ishlasa ham reyting o'zgarmaydi.</i>")
    await safe_edit(cb, text, kb([
        [(f"📖 Darsni o'qish majburiy: {yn(rr)}", "adm:o:require_read")],
        [(f"🔁 Qayta ishlash: {yn(rt)}", "adm:o:allow_retake")],
        [("🏆 Reyting: " + ("birinchi → eng yaxshi" if rm == "first" else "eng yaxshi → birinchi"), "adm:o:rank")],
        [("💾 Zaxira nusxa (baza fayli)", "adm:backup")],
        BACK]))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:o:"))
async def option_toggle(cb: CallbackQuery):
    key = cb.data.split(":")[2]
    if key == "rank":
        await db.set_setting("ranking_mode", "best" if await db.get_setting("ranking_mode") == "first" else "first")
    else:
        await db.set_setting(key, 1 - await db.get_int(key))
    await options(cb)


@router.callback_query(F.data == "adm:backup")
async def backup(cb: CallbackQuery):
    data = await db.dump_json()
    fname = f"zaxira_{datetime.now(TZ).strftime('%Y-%m-%d_%H-%M')}.json"
    await cb.message.answer_document(BufferedInputFile(data, fname),
                                     caption="💾 Bazaning to'liq nusxasi. Xavfsiz joyda saqlang.")
    await cb.answer()


# ======================= Kunlik vazifa =======================
@router.callback_query(F.data == "adm:dly")
async def daily_menu(cb: CallbackQuery):
    from daily import today
    s = {k: await db.get_setting(k) for k in ("daily_on", "daily_time", "daily_count", "daily_group_min",
                                              "remind_on", "remind_time")}
    students = (await db.fetchone("SELECT COUNT(*) AS n FROM users WHERE status='approved'"))["n"]
    done = (await db.fetchone("SELECT COUNT(*) AS n FROM daily_log WHERE day=?", (today(),)))["n"]
    text = ("📅 <b>Kunlik vazifa</b>\n\n"
            f"Holati: {yn(s['daily_on'])} · har kuni <b>{s['daily_time']}</b> da yuboriladi\n"
            f"Savollar soni: <b>{s['daily_count']}</b> ta (har bir bolaga o'z xatolariga qarab tanlanadi)\n"
            f"Guruhda <b>{s['daily_group_min']}</b> ta savolga javob bersa ham bajarilgan hisoblanadi\n"
            f"⏰ Eslatma: {yn(s['remind_on'])} · <b>{s['remind_time']}</b> da (bajarmaganlarga)\n\n"
            f"Bugun bajarganlar: <b>{done} / {students}</b>")
    cnt = int(s["daily_count"])
    gm = int(s["daily_group_min"])
    await safe_edit(cb, text, kb([
        [(("⏸ O'chirish" if int(s["daily_on"]) else "▶️ Yoqish"), "adm:dt:daily_on"),
         ("🕓 Vaqtini o'zgartirish", "adm:dtime:daily_time")],
        [("Savollar soni:", "noop")],
        [(("🔘 " if cnt == n else "") + str(n), f"adm:dset:daily_count:{n}") for n in (5, 10, 15, 20)],
        [("Guruhda nechta javob yetarli:", "noop")],
        [(("🔘 " if gm == n else "") + str(n), f"adm:dset:daily_group_min:{n}") for n in (3, 5, 10)],
        [(f"⏰ Eslatma: {yn(s['remind_on'])}", "adm:dt:remind_on"), ("🕗 Eslatma vaqti", "adm:dtime:remind_time")],
        [("📤 Hozir hammaga yuborish (sinov)", "adm:dnow")],
        BACK]))
    await cb.answer()


@router.callback_query(F.data.startswith("adm:dt:"))
async def daily_toggle(cb: CallbackQuery):
    key = cb.data.split(":")[2]
    await db.set_setting(key, 1 - await db.get_int(key))
    await (parents_menu(cb) if key == "report_on" else daily_menu(cb))


@router.callback_query(F.data.startswith("adm:dset:"))
async def daily_set(cb: CallbackQuery):
    _, _, key, val = cb.data.split(":")
    await db.set_setting(key, val)
    await (parents_menu(cb) if key == "report_day" else daily_menu(cb))


@router.callback_query(F.data.startswith("adm:dtime:"))
async def daily_time_edit(cb: CallbackQuery, state: FSMContext):
    key = cb.data.split(":")[2]
    await state.set_state(AdmSt.time_input)
    await state.update_data(key=key)
    await cb.message.answer(f"🕓 Yangi vaqtni yozing (masalan: <code>16:00</code>). Hozirgi: "
                            f"<b>{await db.get_setting(key)}</b>\nBekor qilish: /cancel")
    await cb.answer()


@router.callback_query(F.data == "adm:dnow")
async def daily_now(cb: CallbackQuery, bot: Bot):
    from daily import send_daily_all
    await cb.answer("Yuborilmoqda…")
    n = await send_daily_all(bot)
    await cb.message.answer(f"📤 Kunlik vazifa {n} ta o'quvchiga yuborildi (bugun bajarganlarga yuborilmaydi).")


# ======================= Ota-onalar =======================
@router.callback_query(F.data == "adm:par")
async def parents_menu(cb: CallbackQuery):
    from parents import WEEKDAYS
    s = {k: await db.get_setting(k) for k in ("report_on", "report_day", "report_time")}
    n_par = (await db.fetchone("SELECT COUNT(DISTINCT parent_id) AS n FROM parents"))["n"]
    students = (await db.fetchone("SELECT COUNT(*) AS n FROM users WHERE status='approved'"))["n"]
    day = int(s["report_day"])
    text = ("👨‍👩‍👦 <b>Ota-onalar</b>\n\n"
            f"Hisobot oladi: <b>{students}</b> o'quvchi akkaunti + <b>{n_par}</b> qo'shimcha telefon\n"
            f"Haftalik hisobot: {yn(s['report_on'])} · har <b>{WEEKDAYS[day]}</b> soat <b>{s['report_time']}</b> da\n\n"
            "<i>Hisobot har bir o'quvchining akkauntiga keladi — bolada telefon bo'lmasa, bu ota-onaning telefoni. "
            "Ikkinchi ota-ona (masalan, dadasi) ham o'z telefonida olishi uchun: botdagi «📊 Haftalik hisobot» "
            "tugmasi → havolani unga yuborish. Ulanganda sizga xabar keladi. O'quvchi kartasida hisobotni ko'rish "
            "va qo'shimcha telefonni uzish mumkin.</i>")
    await safe_edit(cb, text, kb([
        [(("⏸ O'chirish" if int(s["report_on"]) else "▶️ Yoqish"), "adm:dt:report_on"),
         ("🕖 Vaqtini o'zgartirish", "adm:dtime:report_time")],
        [(("🔘 " if day == i else "") + w[:2], f"adm:dset:report_day:{i}") for i, w in enumerate(WEEKDAYS)],
        [("📤 Hisobotni hozir yuborish", "adm:pnow")],
        BACK]))
    await cb.answer()


@router.callback_query(F.data == "adm:pnow")
async def parents_now(cb: CallbackQuery, bot: Bot):
    from parents import send_reports_all
    await cb.answer("Yuborilmoqda…")
    n = await send_reports_all(bot)
    await cb.message.answer(f"📤 Haftalik hisobot {n} ta akkauntga yuborildi.")
